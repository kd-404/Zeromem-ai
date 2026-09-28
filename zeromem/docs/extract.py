"""Get text out of (almost) any document, including scans and photos (OCR).

    extract(path) -> Extracted(pages=[(page_label, text), ...], method, notes)

Formats:
    .pdf                    text layer via pypdf; a page with (almost) no text is a scan, so it is
                            rendered to an image (pypdfium2) and OCR'd
    .png .jpg .jpeg .tif .tiff .bmp .gif .webp .heic    OCR
    .docx                   paragraphs from word/document.xml (no extra package needed)
    .pptx                   one "page" per slide, from ppt/slides/*.xml
    .xlsx                   one line per row: "Header: value, Header: value" (openpyxl)
    .csv .tsv               same row-to-sentence treatment
    .txt .md                as is
    .html .htm              article text (trafilatura), else tags stripped

OCR engines, first one available wins:
    1. Apple Vision via `ocrmac` (macOS only, very good, no model download)
    2. Tesseract via `pytesseract` (any OS; needs the tesseract program: `brew install tesseract`)

Every optional package is imported only when a file needs it, so a missing one breaks only
that file type, with a message saying what to install.
"""

from __future__ import annotations

import csv
import html as htmllib
import io
import re
import zipfile
from dataclasses import dataclass, field
from pathlib import Path

IMAGE_EXT = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".gif", ".webp", ".heic"}
TEXT_EXT = {".txt", ".md", ".markdown"}
SUPPORTED = IMAGE_EXT | TEXT_EXT | {".pdf", ".docx", ".pptx", ".xlsx", ".csv", ".tsv", ".html", ".htm"}
MIN_PAGE_CHARS = 25  # a PDF page with less text than this is treated as a scan and OCR'd


@dataclass
class Extracted:
    pages: list[tuple[str, str]] = field(default_factory=list)  # (label like "3" or "slide 2", text)
    method: str = ""
    notes: list[str] = field(default_factory=list)


# ------------------------------------------------------------------ OCR
def ocr_image(img) -> tuple[str, str]:
    """OCR a PIL image. Returns (text, engine). Raises RuntimeError if no engine is installed."""
    errors = []
    try:  # 1. Apple Vision (macOS)
        from ocrmac import ocrmac
        lines = ocrmac.OCR(img, recognition_level="accurate").recognize()
        return "\n".join(t for t, _conf, _box in lines), "Apple Vision OCR"
    except ImportError:
        errors.append("ocrmac not installed")
    except Exception as e:  # noqa: BLE001 - e.g. not on macOS
        errors.append(f"ocrmac failed: {type(e).__name__}")
    try:  # 2. Tesseract
        import pytesseract
        return pytesseract.image_to_string(img), "Tesseract OCR"
    except ImportError:
        errors.append("pytesseract not installed")
    except Exception as e:  # noqa: BLE001 - usually the tesseract program itself is missing
        errors.append(f"tesseract failed: {type(e).__name__}")
    raise RuntimeError("no OCR engine available (" + "; ".join(errors) + "). On a Mac: pip install ocrmac. "
                       "Anywhere: brew/apt install tesseract, then pip install pytesseract")


def _open_image(path: Path):
    from PIL import Image
    if path.suffix.lower() == ".heic":
        try:
            from pillow_heif import register_heif_opener
            register_heif_opener()
        except ImportError as e:
            raise RuntimeError("HEIC photos need: pip install pillow-heif") from e
    img = Image.open(path)
    return img.convert("RGB")


# ------------------------------------------------------------------ per format
def _pdf(path: Path) -> Extracted:
    from pypdf import PdfReader
    out = Extracted(method="PDF text")
    reader = PdfReader(str(path))
    scanned = []
    for i, page in enumerate(reader.pages, 1):
        text = page.extract_text() or ""
        if len(text.strip()) < MIN_PAGE_CHARS:
            scanned.append(i)
        out.pages.append((str(i), text))
    if scanned:  # render those pages and OCR them
        try:
            import pypdfium2 as pdfium
            pdf = pdfium.PdfDocument(str(path))
            engine = ""
            for i in scanned:
                img = pdf[i - 1].render(scale=2.5).to_pil().convert("RGB")  # ~180 dpi: good for OCR
                text, engine = ocr_image(img)
                out.pages[i - 1] = (str(i), text)
            out.method = f"PDF text + {engine} on {len(scanned)} scanned page(s)" if len(scanned) < len(out.pages) else engine
        except ImportError:
            out.notes.append(f"{len(scanned)} page(s) look scanned but pypdfium2 isn't installed (pip install pypdfium2)")
        except RuntimeError as e:
            out.notes.append(f"{len(scanned)} page(s) look scanned: {e}")
    return out


def _image(path: Path) -> Extracted:
    text, engine = ocr_image(_open_image(path))
    return Extracted(pages=[("1", text)], method=engine)


def _xml_text(xml: str, tag: str) -> list[str]:
    return [htmllib.unescape(t) for t in re.findall(rf"<{tag}(?:\s[^>]*)?>([^<]*)</{tag}>", xml)]


def _docx(path: Path) -> Extracted:
    with zipfile.ZipFile(path) as z:
        xml = z.read("word/document.xml").decode("utf8", "replace")
    paras = []
    for p in re.findall(r"<w:p[ >].*?</w:p>", xml, flags=re.S):
        t = "".join(_xml_text(p, "w:t")).strip()
        if t:
            paras.append(t)
    return Extracted(pages=[("1", "\n".join(paras))], method="Word document")


def _pptx(path: Path) -> Extracted:
    with zipfile.ZipFile(path) as z:
        names = sorted((n for n in z.namelist() if re.fullmatch(r"ppt/slides/slide\d+\.xml", n)),
                       key=lambda n: int(re.search(r"(\d+)", n.rsplit("/", 1)[1]).group(1)))
        pages = []
        for n in names:
            xml = z.read(n).decode("utf8", "replace")
            paras = ["".join(_xml_text(p, "a:t")).strip() for p in re.findall(r"<a:p>.*?</a:p>", xml, flags=re.S)]
            pages.append((f"slide {len(pages) + 1}", "\n".join(x for x in paras if x)))
    return Extracted(pages=pages, method="PowerPoint slides")


def _rows_to_text(rows: list[list[str]]) -> str:
    """Row -> sentence: 'Quarter: Q1, Revenue: 88.' so ZeroMem can point at one row."""
    rows = [[str(c).strip() for c in r] for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return ""
    head, lines = rows[0], []
    for r in rows[1:]:
        cells = [f"{h}: {v}" if h else v for h, v in zip(head + [""] * len(r), r) if v]
        if cells:
            lines.append(", ".join(cells) + ".")
    return "\n".join(lines) if lines else ", ".join(head) + "."


def _xlsx(path: Path) -> Extracted:
    try:
        import openpyxl
    except ImportError as e:
        raise RuntimeError("Excel files need: pip install openpyxl") from e
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pages = [(f"sheet {ws.title}", _rows_to_text([["" if v is None else v for v in r] for r in ws.iter_rows(values_only=True)]))
             for ws in wb.worksheets]
    return Extracted(pages=pages, method="Excel sheets (one sentence per row)")


def _csv(path: Path) -> Extracted:
    raw = path.read_text(encoding="utf8", errors="replace")
    delim = "\t" if path.suffix.lower() == ".tsv" else (csv.Sniffer().sniff(raw[:4096]).delimiter if raw.strip() else ",")
    return Extracted(pages=[("1", _rows_to_text(list(csv.reader(io.StringIO(raw), delimiter=delim))))],
                     method="CSV (one sentence per row)")


def _html(path: Path) -> Extracted:
    raw = path.read_text(encoding="utf8", errors="replace")
    text = None
    try:
        import trafilatura
        text = trafilatura.extract(raw, include_comments=False)
    except ImportError:
        pass
    if not text:
        text = htmllib.unescape(re.sub(r"<[^>]+>", " ", re.sub(r"(?is)<(script|style).*?</\1>", " ", raw)))
    return Extracted(pages=[("1", text)], method="HTML")


def extract(path: str | Path) -> Extracted:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        return _pdf(path)
    if ext in IMAGE_EXT:
        return _image(path)
    if ext == ".docx":
        return _docx(path)
    if ext == ".pptx":
        return _pptx(path)
    if ext == ".xlsx":
        return _xlsx(path)
    if ext in (".csv", ".tsv"):
        return _csv(path)
    if ext in (".html", ".htm"):
        return _html(path)
    if ext in TEXT_EXT:
        return Extracted(pages=[("1", path.read_text(encoding="utf8", errors="replace"))], method="text")
    raise ValueError(f"unsupported file type {ext!r} (supported: {', '.join(sorted(SUPPORTED))})")
